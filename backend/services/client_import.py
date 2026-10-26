"""Load the SGS approved-client master list into the clients table.

    flask --app app import-clients                       # data/clients.json
    flask --app app import-clients --deactivate-others   # and retire the rest

Idempotent: a client whose code or name is already present is skipped, so it
is safe to run again after the list grows. `--deactivate-others` marks every
client that is not on the list and has no CL- code inactive - used once to retire placeholder
clients. They are deactivated rather than deleted, so the jobs and reports that
point at them keep working; they just stop appearing in the pickers.
"""
import json
import os

from extensions import db
from models import Client
from models.client import SBU_OPTIONS, normalise_name

DEFAULT_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "clients.json")


def read_rows(path):
    if path.lower().endswith(".json"):
        with open(path, encoding="utf-8") as handle:
            return json.load(handle)

    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True)["Clients"]
    fields = ("client_code", "client_name", "sbu", "sub_sector")
    return [dict(zip(fields, r[:4])) for r in sheet.iter_rows(min_row=2, values_only=True) if r and r[0] and r[1]]


def import_clients(path=DEFAULT_PATH, deactivate_others=False):
    rows = read_rows(path)
    codes = {code for (code,) in db.session.query(Client.code)}
    names = {normalise_name(name).lower() for (name,) in db.session.query(Client.name)}

    imported = skipped = 0
    listed = set()
    for row in rows:
        code = str(row.get("client_code") or "").strip().upper()
        name = normalise_name(row.get("client_name"))
        if not code or not name:
            continue
        listed.add(code)
        if code in codes or name.lower() in names:
            skipped += 1
            continue
        sbu = (row.get("sbu") or "").strip() or None
        db.session.add(
            Client(
                code=code,
                name=name,
                sbu=sbu if sbu in SBU_OPTIONS else None,
                sub_sector=(row.get("sub_sector") or "").strip() or None,
                is_active=True,
            )
        )
        codes.add(code)
        names.add(name.lower())
        imported += 1

    deactivated = 0
    if deactivate_others:
        for client in Client.query.filter(Client.is_active.is_(True)):
            # CL- codes are real clients - including ones approved after this
            # file was exported - so only other codes are retired.
            if client.code not in listed and not client.code.startswith("CL-"):
                client.is_active = False
                deactivated += 1

    db.session.commit()
    return {"read": len(rows), "imported": imported, "skipped": skipped, "deactivated": deactivated}
